"""Downloadable files (FR13): analysis report (PDF) and claim draft (PDF + editable Word).

PDF: HTML rendered by WeasyPrint, which shapes Arabic and lays out right-to-left through
Pango. The Arabic font (IBM Plex Sans Arabic, SIL OFL) is bundled in backend/assets/fonts.
Word: python-docx with every paragraph marked right-to-left (w:bidi) and complex-script
runs, so the claim opens correctly in Word / Google Docs and can be edited before filing.
"""
import io
from datetime import date
from html import escape
from pathlib import Path

import weasyprint
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.schemas import AnalyzeResponse, ClaimOut

FONTS_DIR = Path(__file__).resolve().parents[1] / "assets" / "fonts"
BLANK = "……………………"

LIKELIHOOD_COLOR = {"مرتفع": "#1f7a4d", "متوسط": "#9a6b00", "منخفض": "#5f5f5a"}

CSS = f"""
@font-face {{ font-family: Plex; src: url("{(FONTS_DIR / 'IBMPlexSansArabic-Regular.ttf').as_uri()}"); font-weight: 400; }}
@font-face {{ font-family: Plex; src: url("{(FONTS_DIR / 'IBMPlexSansArabic-Bold.ttf').as_uri()}"); font-weight: 700; }}
@page {{ size: A4; margin: 18mm 16mm 20mm 16mm;
        @bottom-center {{ content: counter(page) " / " counter(pages); font-family: Plex; font-size: 9pt; color: #777; }} }}
html {{ direction: rtl; }}
body {{ font-family: Plex; font-size: 11pt; line-height: 1.75; color: #1c1c1a; }}
h1 {{ font-size: 17pt; margin: 0 0 4pt; color: #1f5f4a; }}
h2 {{ font-size: 12.5pt; margin: 16pt 0 6pt; padding-bottom: 3pt; border-bottom: 1px solid #ddd; color: #1f5f4a; }}
.meta {{ color: #666; font-size: 9.5pt; }}
.box {{ border: 1px solid #ddd; border-radius: 6pt; padding: 8pt 10pt; margin: 6pt 0; }}
.badge {{ display: inline-block; padding: 1pt 8pt; border-radius: 4pt; color: #fff; font-size: 9.5pt; }}
table {{ width: 100%; border-collapse: collapse; margin: 6pt 0; }}
td, th {{ border: 1px solid #ddd; padding: 5pt 7pt; vertical-align: top; text-align: right; }}
th {{ background: #f1f1ee; font-weight: 700; }}
.num {{ white-space: nowrap; font-weight: 700; }}
.small {{ font-size: 9pt; color: #666; }}
.article {{ font-size: 9.5pt; color: #333; white-space: pre-line; }}
.notes {{ background: #f7f7f5; font-size: 9.5pt; }}
ol, ul {{ margin: 0; padding-right: 18pt; }}
.signature td {{ border: none; padding-top: 14pt; }}
"""


def _sar(amount: float) -> str:
    return f"{amount:,.2f}".rstrip("0").rstrip(".") + " ريال"


def _p(text: str | None) -> str:
    return escape(text or "").replace("\n", "<br>")


def _li(items: list[str], ordered: bool = False) -> str:
    tag = "ol" if ordered else "ul"
    return f"<{tag}>" + "".join(f"<li>{_p(i)}</li>" for i in items) + f"</{tag}>"


def _pdf(body: str, title: str) -> bytes:
    html = f'<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8"><title>{escape(title)}</title><style>{CSS}</style></head><body>{body}</body></html>'
    return weasyprint.HTML(string=html).write_pdf()


def report_pdf(view: AnalyzeResponse) -> bytes:
    a = view.analysis
    o = a.expected_outcome
    parts = [
        f"<h1>تقرير تحليل قضية عمالية</h1>",
        f'<div class="meta">{escape(view.title)} · قضية رقم {view.case_id} · إصدار التحليل {view.version} · '
        f"{view.updated_at:%Y-%m-%d} · قاعدة المعرفة محدّثة حتى {view.knowledge_base_as_of:%Y-%m-%d}</div>",
        "<h2>النتيجة المتوقعة</h2>",
        f'<div class="box"><span class="badge" style="background:{LIKELIHOOD_COLOR[o.likelihood]}">درجة الترجيح: {o.likelihood}</span>'
        f"<p>{_p(o.summary)}</p><p class='small'>السبب: {_p(o.reasoning)}</p></div>",
    ]
    if view.entitlements:
        rows = "".join(
            f"<tr><td>{escape(e.title_ar)}<div class='small'>المادة {escape(e.article)}</div></td>"
            f"<td class='num'>{_sar(e.amount) if e.amount is not None else 'لا يمكن الحساب لنقص معلومات'}</td>"
            f"<td class='small'>{_p(e.formula_ar)}</td></tr>"
            for e in view.entitlements
        )
        total = (
            f"<tr><th>الإجمالي</th><th class='num'>{_sar(view.entitlements_total)}</th><th></th></tr>"
            if view.entitlements_total is not None
            else ""
        )
        parts += [
            "<h2>المستحقات المالية</h2>",
            "<p class='small'>محسوبة آلياً وفق مواد النظام، وليست من تقدير الذكاء الاصطناعي.</p>",
            f"<table><tr><th>المستحق</th><th>المبلغ</th><th>طريقة الحساب</th></tr>{rows}{total}</table>",
        ]
    if a.missing_information:
        parts += ["<h2>معلومات تحتاج لتوضيحها</h2>", _li(a.missing_information)]
    parts += [
        "<h2>ملخص الوقائع</h2>",
        f"<p>{_p(a.facts_summary)}</p>",
        "<h2>المسائل القانونية</h2>",
        _li(a.legal_issues),
        "<h2>التحليل</h2>",
        f"<p>{_p(a.analysis)}</p>",
        "<h2>الخطوات المقترحة</h2>",
        _li(a.recommended_steps, ordered=True),
        "<h2>المواد النظامية المستند إليها</h2>",
    ]
    for c in view.citations:
        parts.append(
            f"<div class='box'><strong>{escape(c.title)}</strong> <span class='small'>— {escape(c.source_name)}</span>"
            f"<p>{_p(c.why)}</p><div class='article'>{escape(c.text)}</div></div>"
        )
    if view.documents:
        parts += ["<h2>المستندات المرفقة</h2>", _li([f"{d.filename} ({'قراءة آلية OCR' if d.method == 'gemini_ocr' else 'نص PDF'})" for d in view.documents])]
    parts += [f"<h2>تنبيه</h2><div class='box notes'>{_p(view.disclaimer)}</div>"]
    return _pdf("".join(parts), f"تقرير قضية {view.case_id}")


def _party_rows(parties) -> str:
    return "".join(f"<tr><th style='width:32%'>{escape(p.label)}</th><td>{escape(p.value) if p.value else BLANK}</td></tr>" for p in parties)


def claim_pdf(claim: ClaimOut) -> bytes:
    plaintiff_name = claim.plaintiff[0].value or BLANK
    parts = [
        "<p class='small'>مسودة صحيفة دعوى عمالية — أُعدّت آلياً للمراجعة</p>",
        f"<p><strong>فضيلة رئيس المحكمة العمالية بـ{escape(claim.court_city) if claim.court_city else BLANK}</strong> حفظه الله</p>",
        "<p>السلام عليكم ورحمة الله وبركاته، وبعد:</p>",
        "<h2>بيانات المدعي</h2>",
        f"<table>{_party_rows(claim.plaintiff)}</table>",
        "<h2>بيانات المدعى عليه</h2>",
        f"<table>{_party_rows(claim.defendant)}</table>",
        f"<h2>موضوع الدعوى</h2><p>{_p(claim.subject)}</p>",
        "<h2>أولاً: الوقائع</h2>",
        _li(claim.facts, ordered=True),
        "<h2>ثانياً: الأسانيد النظامية</h2>",
        _li([f"{g.title}: {g.argument}" for g in claim.legal_grounds], ordered=True),
        "<h2>ثالثاً: الطلبات</h2>",
        "<p>بناءً على ما تقدم، يلتمس المدعي من فضيلتكم الحكم بما يلي:</p>",
    ]
    req_items = []
    for r in claim.requests:
        line = escape(r.text)
        if r.amount is not None:
            line += f" بمبلغ <span class='num'>{_sar(r.amount)}</span><div class='small'>{escape(r.formula or '')}</div>"
        req_items.append(f"<li>{line}</li>")
    parts.append("<ol>" + "".join(req_items) + "</ol>")
    if claim.total is not None:
        parts.append(f"<p><strong>إجمالي المطالبات المالية: <span class='num'>{_sar(claim.total)}</span></strong></p>")
    parts += [
        "<p>وتفضلوا بقبول فائق الاحترام والتقدير،</p>",
        f"<table class='signature'><tr><td>مقدّمه: {escape(plaintiff_name)}</td><td>التوقيع: {BLANK}</td>"
        f"<td>التاريخ: {date.today():%Y/%m/%d} م</td></tr></table>",
        "<h2>ملاحظات</h2>",
        f"<div class='box notes'>{_li(claim.notes)}<p>{_p(claim.disclaimer)}</p></div>",
    ]
    return _pdf("".join(parts), f"مسودة صحيفة دعوى - قضية {claim.case_id}")


# ---------- Word ----------


def _rtl(paragraph, size: float = 12, bold: bool = False, color: RGBColor | None = None):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    ppr.insert(0, bidi)
    for run in paragraph.runs:
        rpr = run._r.get_or_add_rPr()
        rpr.append(OxmlElement("w:rtl"))
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), "Arial")
        cs = OxmlElement("w:szCs")
        cs.set(qn("w:val"), str(int(size * 2)))
        rpr.append(cs)
        if bold:
            rpr.append(OxmlElement("w:bCs"))
        run.font.size = Pt(size)
        run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color
    return paragraph


def _para(doc, text: str, **kw):
    return _rtl(doc.add_paragraph(text), **kw)


def _heading(doc, text: str, color: RGBColor):
    p = _para(doc, text, size=13, bold=True, color=color)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    return p


def _table_rtl(table):
    tbl_pr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tbl_pr.append(bidi)


def claim_docx(claim: ClaimOut) -> bytes:
    doc = Document()
    green = RGBColor(0x1F, 0x5F, 0x4A)
    for section in doc.sections:
        section.right_margin = section.left_margin = Pt(54)
        sect_pr = section._sectPr
        sect_pr.append(OxmlElement("w:bidi"))

    _para(doc, "مسودة صحيفة دعوى عمالية — أُعدّت آلياً للمراجعة", size=9, color=RGBColor(0x66, 0x66, 0x66))
    _para(doc, f"فضيلة رئيس المحكمة العمالية بـ{claim.court_city or BLANK}    حفظه الله", size=13, bold=True)
    _para(doc, "السلام عليكم ورحمة الله وبركاته، وبعد:")

    for heading, parties in (("بيانات المدعي", claim.plaintiff), ("بيانات المدعى عليه", claim.defendant)):
        _heading(doc, heading, green)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        _table_rtl(table)
        for p in parties:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = p.label, p.value or BLANK
            _rtl(cells[0].paragraphs[0], bold=True)
            _rtl(cells[1].paragraphs[0])

    _heading(doc, "موضوع الدعوى", green)
    _para(doc, claim.subject)
    _heading(doc, "أولاً: الوقائع", green)
    for i, fact in enumerate(claim.facts, 1):
        _para(doc, f"{i}. {fact}")
    _heading(doc, "ثانياً: الأسانيد النظامية", green)
    for i, g in enumerate(claim.legal_grounds, 1):
        _para(doc, f"{i}. {g.title}: {g.argument}")
    _heading(doc, "ثالثاً: الطلبات", green)
    _para(doc, "بناءً على ما تقدم، يلتمس المدعي من فضيلتكم الحكم بما يلي:")
    for i, r in enumerate(claim.requests, 1):
        _para(doc, f"{i}. {r.text}" + (f" بمبلغ {_sar(r.amount)}" if r.amount is not None else ""))
        if r.formula:
            _para(doc, f"   ({r.formula})", size=9, color=RGBColor(0x66, 0x66, 0x66))
    if claim.total is not None:
        _para(doc, f"إجمالي المطالبات المالية: {_sar(claim.total)}", bold=True)
    _para(doc, "وتفضلوا بقبول فائق الاحترام والتقدير،")
    _para(doc, f"مقدّمه: {claim.plaintiff[0].value or BLANK}      التوقيع: {BLANK}      التاريخ: {date.today():%Y/%m/%d} م")

    _para(doc, "ملاحظات", size=11, bold=True, color=green)
    for note in claim.notes + [claim.disclaimer]:
        _para(doc, f"• {note}", size=9, color=RGBColor(0x66, 0x66, 0x66))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
