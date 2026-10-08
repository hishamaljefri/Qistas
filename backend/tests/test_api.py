"""API tests on the qistas_test database, with Gemini faked (no API calls, no quota used)."""
import io
from datetime import UTC, datetime, timedelta
from pathlib import Path

import jwt
import pymupdf
from docx import Document
from sqlalchemy import text

from app.config import JWT_SECRET
from app.db import engine
from tests.conftest import CASE, register

FIXTURES = Path(__file__).parent / "fixtures"


def make_admin(username: str) -> None:
    with engine.begin() as conn:
        conn.execute(text("UPDATE users SET role='admin' WHERE username=:u"), {"u": username})


# ---------- accounts (FR1, FR2, FR14; SR1, SR2, SR4) ----------


def test_register_login_me_refresh(client):
    register(client, "Khaled")  # stored lowercase
    r = client.post("/api/auth/login", json={"identifier": "KHALED@example.com", "password": "password123"})
    assert r.status_code == 200
    token = r.json()["access_token"]
    me = client.get("/api/auth/me", headers={"Authorization": f"Bearer {token}"}).json()
    assert me["username"] == "khaled" and me["role"] == "user"
    assert client.post("/api/auth/refresh", headers={"Authorization": f"Bearer {token}"}).status_code == 200


def test_password_is_stored_as_bcrypt_hash(client):
    register(client, "sara", password="my-secret-pass")
    with engine.connect() as conn:
        stored = conn.execute(text("SELECT password_hash FROM users WHERE username='sara'")).scalar()
    assert stored.startswith("$2b$") and "my-secret-pass" not in stored


def test_duplicate_and_wrong_password(client):
    register(client, "khaled")
    dup = client.post("/api/auth/register", json={"username": "khaled", "email": "x@example.com", "password": "password123"})
    assert dup.status_code == 409
    bad = client.post("/api/auth/login", json={"identifier": "khaled", "password": "wrong-password"})
    assert bad.status_code == 401
    unknown = client.post("/api/auth/login", json={"identifier": "nobody", "password": "password123"})
    assert unknown.status_code == 401 and unknown.json()["detail"] == bad.json()["detail"]


def test_weak_input_rejected(client):
    r = client.post("/api/auth/register", json={"username": "ab", "email": "bad", "password": "short"})
    assert r.status_code == 422


def test_expired_and_forged_tokens_rejected(client):
    register(client, "khaled")
    past = datetime.now(UTC) - timedelta(minutes=31)
    expired = jwt.encode({"sub": "1", "role": "user", "iat": past, "exp": past + timedelta(minutes=30)}, JWT_SECRET)
    r = client.get("/api/auth/me", headers={"Authorization": f"Bearer {expired}"})
    assert r.status_code == 401 and "انتهت الجلسة" in r.json()["detail"]
    forged = jwt.encode({"sub": "1", "role": "admin", "exp": datetime.now(UTC) + timedelta(minutes=5)}, "not-the-secret")
    assert client.get("/api/auth/me", headers={"Authorization": f"Bearer {forged}"}).status_code == 401


def test_admin_manages_users_and_roles(client):
    admin_h = register(client, "admin1")
    make_admin("admin1")
    user_h = register(client, "khaled")
    assert client.get("/api/admin/users", headers=user_h).status_code == 403  # SR4
    users = client.get("/api/admin/users", headers=admin_h).json()
    khaled = next(u for u in users if u["username"] == "khaled")
    # deactivate: the user's existing token stops working and login is refused
    assert client.patch(f"/api/admin/users/{khaled['id']}", json={"is_active": False}, headers=admin_h).status_code == 200
    assert client.get("/api/auth/me", headers=user_h).status_code == 401
    assert client.post("/api/auth/login", json={"identifier": "khaled", "password": "password123"}).status_code == 403
    # an admin cannot remove their own admin role
    me = next(u for u in users if u["username"] == "admin1")
    assert client.patch(f"/api/admin/users/{me['id']}", json={"role": "user"}, headers=admin_h).status_code == 400


# ---------- cases, privacy, ownership (FR3, FR12; SR7, SR8, SR9) ----------


def test_analyze_requires_login(client, gemini):
    assert client.post("/api/cases/analyze", json=CASE).status_code == 401
    assert gemini.prompts == []


def test_real_values_never_reach_gemini_or_database_but_owner_sees_them(client, gemini):
    h = register(client, "khaled")
    r = client.post("/api/cases/analyze", json=CASE, headers=h)
    assert r.status_code == 201, r.text
    body = r.json()
    secrets = ["خالد", "الغامدي", "1087654321", "0551234567", "الريادة"]

    sent = " ".join(gemini.prompts)
    assert all(s not in sent for s in secrets) and "[EMPLOYEE_1]" in sent  # SR7/SR8

    with engine.connect() as conn:
        stored = conn.execute(
            text("SELECT string_agg(t, ' ') FROM (SELECT c::text AS t FROM cases c UNION ALL SELECT a::text FROM case_analyses a) x")
        ).scalar()
    assert all(s not in stored for s in secrets)  # SR9: only masked text + encrypted vault in the DB

    assert "خالد سعد الغامدي" in body["description"] and "خالد سعد الغامدي" in body["title"]
    assert body["entitlements_total"] == 65000  # 35,000 end-of-service + 30,000 art. 77, computed in code
    assert [c["record_id"] for c in body["citations"]] == ["LAW-0077", "LAW-0084"]


def test_my_cases_list_rename_reanalyze_delete(client, gemini):
    h = register(client, "khaled")
    case_id = client.post("/api/cases/analyze", json=CASE, headers=h).json()["case_id"]

    cases = client.get("/api/cases", headers=h).json()
    assert len(cases) == 1 and cases[0]["likelihood"] == "مرتفع" and cases[0]["entitlements_total"] == 65000

    renamed = client.patch(f"/api/cases/{case_id}", json={"title": "قضية خالد الغامدي ضد الريادة"}, headers=h).json()
    # short forms share the full name's placeholder, so they come back as the full name
    assert renamed["title"] == "قضية خالد سعد الغامدي ضد الريادة للمقاولات"
    with engine.connect() as conn:
        stored_title = conn.execute(text("SELECT title FROM cases WHERE id=:i"), {"i": case_id}).scalar()
    assert "الغامدي" not in stored_title  # titles are masked too

    edited = client.post(
        f"/api/cases/{case_id}/reanalyze",
        json={"description": CASE["description"] + " وأضيف أن راتبي 10000 ريال."},
        headers=h,
    ).json()
    assert edited["version"] == 2 and "خالد سعد الغامدي" in edited["description"]

    assert client.delete(f"/api/cases/{case_id}", headers=h).status_code == 204
    assert client.get(f"/api/cases/{case_id}", headers=h).status_code == 404


def test_users_cannot_reach_each_others_cases(client, gemini):
    owner = register(client, "khaled")
    other = register(client, "sara")
    case_id = client.post("/api/cases/analyze", json=CASE, headers=owner).json()["case_id"]
    for method, path in [("get", ""), ("delete", ""), ("get", "/report.pdf"), ("post", "/claim")]:
        kwargs = {"json": {}} if method == "post" else {}
        r = getattr(client, method)(f"/api/cases/{case_id}{path}", headers=other, **kwargs)
        assert r.status_code == 404, (method, path)
    assert client.get("/api/cases", headers=other).json() == []
    admin = register(client, "admin1")
    make_admin("admin1")
    assert client.get(f"/api/cases/{case_id}", headers=admin).status_code == 200


# ---------- documents (FR4, FR5; SR6) ----------


def _upload(client, h, name, data, **form):
    return client.post("/api/documents/extract", files={"file": (name, data)}, data={k: str(v).lower() for k, v in form.items()}, headers=h)


def test_text_pdf_is_read_locally(client, gemini):
    h = register(client, "khaled")
    r = _upload(client, h, "regulations.pdf", (FIXTURES / "hrsd_regulations_page4.pdf").read_bytes())
    assert r.status_code == 200, r.text
    assert r.json()["method"] == "text_layer" and "المنشأة" in r.json()["text"]
    assert gemini.prompts == []  # nothing sent to Gemini


def test_scan_needs_consent_then_uses_gemini(client, gemini):
    h = register(client, "khaled")
    png = (FIXTURES / "scanned_page.png").read_bytes()
    r = _upload(client, h, "scan.png", png)
    assert r.status_code == 422 and r.json()["detail"]["code"] == "consent_required"
    assert gemini.prompts == []
    r = _upload(client, h, "scan.png", png, cloud_ocr_consent=True)
    assert r.status_code == 200 and r.json()["method"] == "gemini_ocr"


def test_unsafe_or_oversized_files_rejected(client, gemini):
    h = register(client, "khaled")
    assert _upload(client, h, "invoice.pdf", b"MZ\x90\x00" + b"\x00" * 100).status_code == 415  # an .exe renamed .pdf
    assert _upload(client, h, "big.pdf", b"%PDF-" + b"0" * (10 * 1024 * 1024)).status_code == 413
    assert client.post("/api/documents/extract", files={"file": ("a.pdf", b"%PDF-")}).status_code == 401


def test_documents_are_masked_and_attached_to_the_case(client, gemini):
    h = register(client, "khaled")
    doc = {"filename": "عقد خالد.pdf", "method": "text_layer", "pages": 1, "text": "عقد عمل بين شركة الريادة للمقاولات وخالد سعد الغامدي هوية 1087654321"}
    body = client.post("/api/cases/analyze", json={**CASE, "documents": [doc]}, headers=h).json()
    assert body["documents"][0]["filename"] == "عقد خالد سعد الغامدي.pdf"  # short name restored as the full name
    assert "مستند مرفق" in gemini.prompts[-1] and "1087654321" not in gemini.prompts[-1]
    with engine.connect() as conn:
        stored = conn.execute(text("SELECT filename || masked_text FROM case_documents")).scalar()
    assert "الغامدي" not in stored and "1087654321" not in stored


# ---------- claim (FR10) and downloads (FR13) ----------


def test_claim_draft_and_downloads(client, gemini):
    h = register(client, "khaled")
    case_id = client.post("/api/cases/analyze", json=CASE, headers=h).json()["case_id"]
    assert client.get(f"/api/cases/{case_id}/claim", headers=h).status_code == 404

    party = {"plaintiff_nationality": "سعودي", "defendant_cr_number": "4030123456", "court_city": "جدة"}
    claim = client.post(f"/api/cases/{case_id}/claim", json=party, headers=h).json()
    assert claim["plaintiff"][0]["value"] == "خالد سعد الغامدي"
    assert claim["plaintiff"][1]["value"] == "1087654321"  # taken from the vault, never from Gemini
    assert "4030123456" not in " ".join(gemini.prompts) and "سعودي" not in gemini.prompts[-1]
    assert [g["record_id"] for g in claim["legal_grounds"]] == ["LAW-0077"]  # invented LAW-9999 removed
    assert claim["grounding_warnings"] and "LAW-9999" in claim["grounding_warnings"][0]
    assert [r["amount"] for r in claim["requests"]] == [35000, 30000, None] and claim["total"] == 65000

    report = client.get(f"/api/cases/{case_id}/report.pdf", headers=h)
    assert report.status_code == 200 and report.headers["content-type"] == "application/pdf"
    assert pymupdf.open(stream=report.content, filetype="pdf").page_count >= 1

    pdf = client.get(f"/api/cases/{case_id}/claim.pdf", headers=h)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")

    docx = client.get(f"/api/cases/{case_id}/claim.docx", headers=h)
    assert docx.status_code == 200 and "attachment" in docx.headers["content-disposition"]
    word = Document(io.BytesIO(docx.content))
    all_text = "\n".join(p.text for p in word.paragraphs) + "".join(c.text for t in word.tables for r in t.rows for c in r.cells)
    assert "خالد سعد الغامدي" in all_text and "35,000 ريال" in all_text
    assert "w:bidi" in word.paragraphs[0]._p.xml  # right-to-left paragraphs


def test_published_text_refusal_gives_a_clear_message(client, gemini, monkeypatch):
    from app import documents
    from app.gemini import GenerationBlocked

    def blocked(*_, **__):
        raise GenerationBlocked("RECITATION")

    monkeypatch.setattr(documents, "generate", blocked)
    h = register(client, "khaled")
    r = _upload(client, h, "law.png", (FIXTURES / "scanned_page.png").read_bytes(), cloud_ocr_consent=True)
    assert r.status_code == 422 and r.json()["detail"]["code"] == "published_text"
